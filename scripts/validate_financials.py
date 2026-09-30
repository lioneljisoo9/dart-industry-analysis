import argparse
import json
import re
import sys
import unicodedata
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from config.companies import COMPANIES  # noqa: E402
from src.storage import (  # noqa: E402
    MASTER_DIRECTORY,
    company_storage_key,
    load_company_storage,
    storage_exists,
)
from src.validation import (  # noqa: E402
    build_financial_validation_report,
    validation_summary,
)


PROCESSED_DIRECTORY = PROJECT_ROOT / "data" / "processed"
OUTPUT_DIRECTORY = PROJECT_ROOT / "outputs" / "excel"


def _display_width(value):
    return sum(
        2
        if unicodedata.east_asian_width(character)
        in {"W", "F"}
        else 1
        for character in str(value)
    )


def _normalize_query(value):
    return re.sub(
        r"[\s\-_.·,]+",
        "",
        str(value or "").casefold(),
    )


def _metadata_targets():
    targets = []

    for path in MASTER_DIRECTORY.glob("*_metadata.json"):
        payload = json.loads(
            path.read_text(encoding="utf-8")
        )
        targets.append(
            {
                "company_name": str(
                    payload.get("company_name") or ""
                ).strip(),
                "stock_code": str(
                    payload.get("stock_code") or ""
                ).strip(),
                "corp_code": str(
                    payload.get("corp_code") or ""
                ).strip(),
            }
        )

    return targets


def _configured_targets():
    return [
        {
            "company_name": str(config["corp_name"]),
            "stock_code": str(
                config.get("stock_code") or ""
            ).strip(),
            "corp_code": "",
        }
        for config in COMPANIES.values()
    ]


def _merge_targets(*target_groups):
    merged = {}

    for targets in target_groups:
        for target in targets:
            key = company_storage_key(
                stock_code=target.get("stock_code"),
                corp_code=target.get("corp_code"),
            )
            existing = merged.get(key, {})
            merged[key] = {
                "company_name": (
                    target.get("company_name")
                    or existing.get("company_name")
                    or key
                ),
                "stock_code": (
                    target.get("stock_code")
                    or existing.get("stock_code")
                    or ""
                ),
                "corp_code": (
                    target.get("corp_code")
                    or existing.get("corp_code")
                    or ""
                ),
                "storage_key": key,
            }

    return list(merged.values())


def _select_targets(queries):
    available = _merge_targets(
        _configured_targets(),
        _metadata_targets(),
    )

    if not queries:
        return [
            target
            for target in available
            if target["company_name"]
            in {
                config["corp_name"]
                for config in COMPANIES.values()
            }
        ]

    selected = []

    for query in queries:
        query_key = _normalize_query(query)
        matches = [
            target
            for target in available
            if query_key
            in {
                _normalize_query(target["company_name"]),
                _normalize_query(target["stock_code"]),
                _normalize_query(target["corp_code"]),
                _normalize_query(target["storage_key"]),
            }
        ]

        if not matches:
            raise ValueError(
                f"검증할 로컬 Master를 찾지 못했습니다: {query}"
            )
        if len(matches) > 1:
            names = ", ".join(
                target["company_name"]
                for target in matches
            )
            raise ValueError(
                f"검증 대상이 여러 개입니다: {query} → {names}"
            )

        if matches[0] not in selected:
            selected.append(matches[0])

    return selected


def _style_validation_workbook(writer):
    header_fill = PatternFill(
        "solid",
        fgColor="1F4E78",
    )
    status_fills = {
        "PASS_EXACT": PatternFill(
            "solid",
            fgColor="D9EAD3",
        ),
        "WARN_ROUNDING": PatternFill(
            "solid",
            fgColor="FFF2CC",
        ),
        "WARN_ACCOUNT_ID_CHANGED": PatternFill(
            "solid",
            fgColor="FFF2CC",
        ),
    }
    failure_fill = PatternFill(
        "solid",
        fgColor="F4CCCC",
    )
    number_headers = {
        "annual_amount",
        "comparison_amount",
        "difference",
        "absolute_difference",
        "tolerance",
    }

    for worksheet in writer.book.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.sheet_view.showGridLines = False

        header_map = {
            cell.value: cell.column
            for cell in worksheet[1]
            if cell.value is not None
        }

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = Font(
                bold=True,
                color="FFFFFF",
            )

        for header in number_headers:
            column = header_map.get(header)
            if not column:
                continue
            for row in range(2, worksheet.max_row + 1):
                worksheet.cell(
                    row=row,
                    column=column,
                ).number_format = "#,##0"

        status_column = header_map.get("validation_status")
        hard_failure_column = header_map.get("hard_failure")

        if status_column:
            for row in range(2, worksheet.max_row + 1):
                status_cell = worksheet.cell(
                    row=row,
                    column=status_column,
                )
                hard_failure = (
                    worksheet.cell(
                        row=row,
                        column=hard_failure_column,
                    ).value
                    if hard_failure_column
                    else False
                )
                status_cell.fill = (
                    failure_fill
                    if hard_failure
                    else status_fills.get(
                        status_cell.value,
                        PatternFill(
                            "solid",
                            fgColor="EDEDED",
                        ),
                    )
                )

        for column_cells in worksheet.columns:
            max_length = max(
                (
                    _display_width(cell.value)
                    for cell in column_cells
                    if cell.value is not None
                ),
                default=0,
            )
            column_letter = get_column_letter(
                column_cells[0].column
            )
            worksheet.column_dimensions[
                column_letter
            ].width = min(max_length + 2, 45)


def _write_validation_workbook(report_df, path):
    status_summary = (
        report_df.pivot_table(
            index="company",
            columns="validation_status",
            aggfunc="size",
            fill_value=0,
        )
        .reset_index()
    )
    status_summary.columns.name = None
    status_summary.insert(
        1,
        "total_rules",
        report_df.groupby("company").size().reindex(
            status_summary["company"]
        ).to_numpy(),
    )
    status_summary.insert(
        2,
        "hard_failures",
        report_df.groupby("company")["hard_failure"]
        .sum()
        .reindex(status_summary["company"])
        .astype(int)
        .to_numpy(),
    )
    status_summary = status_summary.sort_values("company")

    path.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    with pd.ExcelWriter(
        path,
        engine="openpyxl",
    ) as writer:
        status_summary.to_excel(
            writer,
            sheet_name="Summary",
            index=False,
        )
        report_df.to_excel(
            writer,
            sheet_name="Validation",
            index=False,
        )
        _style_validation_workbook(writer)


def parse_args():
    parser = argparse.ArgumentParser(
        description=(
            "로컬 Parquet Master의 연간/분기 재무데이터를 "
            "회계 규칙으로 검증합니다."
        )
    )
    parser.add_argument(
        "--company",
        action="append",
        dest="companies",
        help=(
            "검증할 기업명·종목코드·corp_code. "
            "여러 기업은 옵션을 반복합니다."
        ),
    )
    parser.add_argument(
        "--tolerance",
        type=int,
        default=1000,
        help="반올림 경고로 허용할 절대 금액 차이(기본 1,000원)",
    )
    parser.add_argument(
        "--allow-failures",
        action="store_true",
        help="hard failure가 있어도 종료코드 0을 반환합니다.",
    )
    return parser.parse_args()


def main():
    args = parse_args()
    targets = _select_targets(args.companies)

    if not targets:
        raise ValueError("검증할 기업이 없습니다.")

    reports = []
    PROCESSED_DIRECTORY.mkdir(
        parents=True,
        exist_ok=True,
    )

    for target in targets:
        company_name = target["company_name"]
        stock_code = target["stock_code"]
        corp_code = target["corp_code"]
        storage_key = target["storage_key"]

        if not storage_exists(
            stock_code=stock_code,
            corp_code=corp_code,
        ):
            raise FileNotFoundError(
                f"{company_name}의 Master Data가 없습니다. "
                "먼저 main.py 또는 scripts/update_financials.py를 "
                "실행하세요."
            )

        annual_df, quarterly_df = load_company_storage(
            stock_code=stock_code,
            corp_code=corp_code,
        )
        report = build_financial_validation_report(
            annual_df=annual_df,
            quarterly_df=quarterly_df,
            rounding_tolerance=args.tolerance,
        )
        report.insert(1, "storage_key", storage_key)
        reports.append(report)

        company_path = (
            PROCESSED_DIRECTORY
            / f"{storage_key}_validation.parquet"
        )
        report.to_parquet(
            company_path,
            index=False,
        )

        summary = validation_summary(report)
        print()
        print(f"=== {company_name} 검증 결과 ===")
        print(f"전체 규칙: {summary['total']}")
        print(f"Hard failure: {summary['hard_failures']}")
        for status, count in summary["status_counts"].items():
            print(f"  {status}: {count}")

    combined = pd.concat(
        reports,
        ignore_index=True,
    )
    combined_path = (
        PROCESSED_DIRECTORY
        / "financial_validation.parquet"
    )
    combined.to_parquet(
        combined_path,
        index=False,
    )

    excel_path = (
        OUTPUT_DIRECTORY
        / "financial_validation.xlsx"
    )
    _write_validation_workbook(
        combined,
        excel_path,
    )

    combined_summary = validation_summary(combined)

    print()
    print(f"통합 Parquet: {combined_path}")
    print(f"검증 Excel: {excel_path}")
    print(
        "통합 Hard failure: "
        f"{combined_summary['hard_failures']}"
    )

    if (
        combined_summary["hard_failures"]
        and not args.allow_failures
    ):
        raise SystemExit(1)


if __name__ == "__main__":
    main()

