from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from src.normalize import extract_standard_accounts


def _reorder_columns(df):
    """
    자주 확인할 컬럼을 앞쪽에 배치하고,
    나머지 DART 원본 컬럼은 뒤에 그대로 보존한다.
    """

    preferred_columns = [
        "company",
        "year",
        "quarter",
        "quarter_num",
        "period",
        "fs_div_source",
        "rcept_no",
        "reprt_code",
        "bsns_year",
        "sj_div",
        "sj_nm",
        "account_id",
        "account_nm",
        "account_detail",
        "quarter_amount",
        "quarter_amount_basis",
        "calculation_status",
        "thstrm_nm",
        "thstrm_amount",
        "thstrm_add_amount",
        "frmtrm_nm",
        "frmtrm_amount",
        "frmtrm_q_nm",
        "frmtrm_q_amount",
        "frmtrm_add_amount",
        "bfefrmtrm_nm",
        "bfefrmtrm_amount",
        "ord",
        "currency",
    ]

    existing_preferred = [
        col for col in preferred_columns
        if col in df.columns
    ]

    remaining_columns = [
        col for col in df.columns
        if col not in existing_preferred
    ]

    return df[existing_preferred + remaining_columns].copy()


def _format_worksheet(ws):
    """Excel 시트 기본 서식 설정."""

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    for cell in ws[1]:
        cell.font = Font(bold=True)

    amount_columns = {
        "thstrm_amount",
        "thstrm_add_amount",
        "frmtrm_amount",
        "frmtrm_q_amount",
        "frmtrm_add_amount",
        "bfefrmtrm_amount",
        "amount",
        "quarter_amount",
    }

    header_map = {
        cell.value: cell.column
        for cell in ws[1]
        if cell.value is not None
    }

    for column_name in amount_columns:
        if column_name not in header_map:
            continue

        column_number = header_map[column_name]

        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=column_number)

            if isinstance(cell.value, (int, float)):
                cell.number_format = "#,##0"

    for column_cells in ws.columns:
        max_length = 0
        column_letter = get_column_letter(column_cells[0].column)

        for cell in column_cells:
            if cell.value is None:
                continue
            max_length = max(max_length, len(str(cell.value)))

        ws.column_dimensions[column_letter].width = min(
            max_length + 2,
            35,
        )


def _build_core_dataframe(df, quarterly=False):
    core_df = extract_standard_accounts(df)

    if core_df.empty:
        return core_df

    columns = [
        "company",
        "year",
    ]

    if quarterly:
        columns.extend([
            "quarter",
            "period",
        ])
    else:
        columns.append("period")

    columns.extend([
        "sj_div",
        "account_id",
        "account_nm",
        "standard_account",
        "standard_account_match",
    ])

    if quarterly:
        columns.extend([
            "quarter_amount",
            "quarter_amount_basis",
            "calculation_status",
            "thstrm_amount",
            "thstrm_add_amount",
        ])
    else:
        columns.append("thstrm_amount")

    columns = [col for col in columns if col in core_df.columns]
    return core_df[columns].copy()


def _fallback_excel_path(file_path):
    """열려 있는 Excel 파일을 덮어쓸 수 없을 때 사용할 경로."""

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )
    candidate = file_path.with_name(
        f"{file_path.stem}_{timestamp}"
        f"{file_path.suffix}"
    )
    counter = 1

    while candidate.exists():
        candidate = file_path.with_name(
            f"{file_path.stem}_{timestamp}_"
            f"{counter}{file_path.suffix}"
        )
        counter += 1

    return candidate


def _write_financial_workbook(
    file_path,
    annual_all,
    annual_core,
    quarterly_all,
    quarterly_core,
):
    with pd.ExcelWriter(
        file_path,
        engine="openpyxl",
    ) as writer:
        if not annual_all.empty:
            annual_all.to_excel(
                writer,
                sheet_name="Annual_All",
                index=False,
            )

        if not annual_core.empty:
            annual_core.to_excel(
                writer,
                sheet_name="Annual_Core",
                index=False,
            )

        if not quarterly_all.empty:
            quarterly_all.to_excel(
                writer,
                sheet_name="Quarterly_All",
                index=False,
            )

        if not quarterly_core.empty:
            quarterly_core.to_excel(
                writer,
                sheet_name="Quarterly_Core",
                index=False,
            )

        for worksheet in writer.book.worksheets:
            _format_worksheet(worksheet)


def export_financial_workbook(
    annual_df,
    quarterly_df,
    file_path,
):
    """
    최근 연간/분기 전체계정과 Core 계정을 하나의 Workbook으로 저장한다.
    """

    if annual_df.empty and quarterly_df.empty:
        raise ValueError("저장할 재무데이터가 없습니다.")

    file_path = Path(file_path)
    file_path.parent.mkdir(parents=True, exist_ok=True)

    annual_all = _reorder_columns(annual_df) if not annual_df.empty else pd.DataFrame()
    quarterly_all = _reorder_columns(quarterly_df) if not quarterly_df.empty else pd.DataFrame()

    annual_core = (
        _build_core_dataframe(annual_df, quarterly=False)
        if not annual_df.empty
        else pd.DataFrame()
    )
    quarterly_core = (
        _build_core_dataframe(quarterly_df, quarterly=True)
        if not quarterly_df.empty
        else pd.DataFrame()
    )

    actual_path = file_path

    try:
        _write_financial_workbook(
            file_path=actual_path,
            annual_all=annual_all,
            annual_core=annual_core,
            quarterly_all=quarterly_all,
            quarterly_core=quarterly_core,
        )
    except PermissionError:
        actual_path = _fallback_excel_path(
            file_path
        )
        print()
        print(
            "기존 Excel 파일이 열려 있어 덮어쓸 수 없습니다. "
            "새 파일로 저장합니다: "
            f"{actual_path}"
        )
        _write_financial_workbook(
            file_path=actual_path,
            annual_all=annual_all,
            annual_core=annual_core,
            quarterly_all=quarterly_all,
            quarterly_core=quarterly_core,
        )

    print()
    print(f"Excel 저장 완료: {actual_path}")
    print(f"Annual_All: {len(annual_all):,}행")
    print(f"Annual_Core: {len(annual_core):,}행")
    print(f"Quarterly_All: {len(quarterly_all):,}행")
    print(f"Quarterly_Core: {len(quarterly_core):,}행")

    return actual_path


def export_annual_workbook(
    annual_df,
    file_path,
):
    """기존 호출과의 호환을 위한 연간 전용 wrapper."""

    export_financial_workbook(
        annual_df=annual_df,
        quarterly_df=pd.DataFrame(),
        file_path=file_path,
    )

